"""Регрессионные тесты autofix на эталонном документе кафедры.

Эталон (tests/fixtures/reference_agu.docx) принят нормоконтролем без
замечаний. Главный инвариант: autofix не должен ломать то, что уже
правильно. Проверяем структуру, текст, титульный лист, колонтитулы,
разделы, а также что после autofix число замечаний не растёт.
"""
from __future__ import annotations

import io
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Cm, Pt

from app.checkers.docx_checker import check_docx_document, _find_body_start
from app.checkers.docx_fixer import autofix_docx

REFERENCE = Path(__file__).parent / "fixtures" / "reference_agu.docx"


@pytest.fixture(scope="module")
def fixed_pair(tmp_path_factory):
    tmp = tmp_path_factory.mktemp("ref")
    fixed_path = tmp / "reference_fixed.docx"
    fixed_path.write_bytes(autofix_docx(REFERENCE))
    return Document(REFERENCE), Document(fixed_path), fixed_path


def _pf(p):
    f = p.paragraph_format
    return (p.alignment, f.first_line_indent, f.left_indent, f.line_spacing,
            f.space_before, f.space_after)


def test_structure_preserved(fixed_pair):
    ref, fixed, _ = fixed_pair
    assert len(fixed.paragraphs) == len(ref.paragraphs), "autofix удалил или добавил абзацы"
    assert len(fixed.sections) == len(ref.sections), "autofix потерял раздел документа"
    assert len(fixed.tables) == len(ref.tables)


def test_text_unchanged(fixed_pair):
    ref, fixed, _ = fixed_pair
    assert [p.text for p in fixed.paragraphs] == [p.text for p in ref.paragraphs]
    for rt, ft in zip(ref.tables, fixed.tables):
        for rr, fr in zip(rt.rows, ft.rows):
            for rc, fc in zip(rr.cells, fr.cells):
                assert rc.text == fc.text


def test_title_page_untouched(fixed_pair):
    """Всё до начала основного текста — как в эталоне (поля не в счёт)."""
    ref, fixed, _ = fixed_pair
    body_start = _find_body_start(ref)
    assert body_start > 0
    for i in range(body_start):
        assert _pf(fixed.paragraphs[i]) == _pf(ref.paragraphs[i]), f"титульный лист: абзац {i} изменён"
        for rr, fr in zip(ref.paragraphs[i].runs, fixed.paragraphs[i].runs):
            assert (rr.bold, rr.italic, rr.font.size, rr.font.name) == (fr.bold, fr.italic, fr.font.size, fr.font.name)


def test_headers_footers_and_sections_preserved(fixed_pair):
    ref, fixed, _ = fixed_pair
    for rs, fs in zip(ref.sections, fixed.sections):
        assert rs.start_type == fs.start_type
        assert rs.orientation == fs.orientation
        assert rs.different_first_page_header_footer == fs.different_first_page_header_footer
        assert [p.text for p in rs.header.paragraphs] == [p.text for p in fs.header.paragraphs]
        assert [p.text for p in rs.footer.paragraphs] == [p.text for p in fs.footer.paragraphs]
        assert rs.footer._element.xml == fs.footer._element.xml


def test_section_break_paragraphs_kept(fixed_pair):
    ref, fixed, _ = fixed_pair
    count = lambda d: sum(1 for p in d.paragraphs if p._element.pPr is not None
                          and p._element.pPr.find(qn("w:sectPr")) is not None)
    assert count(fixed) == count(ref)


def test_heading_text_with_multiple_spaces_untouched(fixed_pair):
    """«1.1   Название» — пробелы после номера намеренные."""
    ref, fixed, _ = fixed_pair
    for rp, fp in zip(ref.paragraphs, fixed.paragraphs):
        if "подраздел" in (rp.style.name or "").lower():
            assert rp.text == fp.text


def test_centered_body_paragraphs_not_justified(fixed_pair):
    ref, fixed, _ = fixed_pair
    for rp, fp in zip(ref.paragraphs, fixed.paragraphs):
        if rp.alignment == WD_ALIGN_PARAGRAPH.CENTER:
            assert fp.alignment == WD_ALIGN_PARAGRAPH.CENTER


def test_font_size_12_not_forced_to_14(fixed_pair):
    """В эталоне основной текст 12 пт — это допустимый размер, не трогаем."""
    ref, fixed, _ = fixed_pair
    assert fixed.styles["Normal"].font.size == ref.styles["Normal"].font.size


def test_no_explicit_formatting_spam(fixed_pair):
    """autofix не должен размазывать явный шрифт/цвет по всем run-ам."""
    ref, fixed, _ = fixed_pair
    def explicit(d):
        return sum(1 for p in d.paragraphs for r in p.runs
                   if r._element.rPr is not None and r._element.rPr.find(qn("w:rFonts")) is not None)
    assert explicit(fixed) <= explicit(ref) + 10


def test_issues_do_not_increase_and_verdict_not_bad(fixed_pair):
    ref, fixed, fixed_path = fixed_pair
    before = check_docx_document(REFERENCE, "ref.docx")
    after = check_docx_document(fixed_path, "fixed.docx")
    assert after["total_issues"] <= before["total_issues"]
    assert after["summary"]["high"] == 0
    assert after["verdict"] in ("good", "ok")
    assert after["total_issues"] <= 3, [i["code"] for i in after["issues"]]


def test_autofix_is_idempotent(fixed_pair):
    """Повторный autofix ничего не меняет в уже исправленном документе."""
    _, _, fixed_path = fixed_pair
    twice = Document(io.BytesIO(autofix_docx(fixed_path)))
    once = Document(fixed_path)
    assert [p.text for p in twice.paragraphs] == [p.text for p in once.paragraphs]
    assert [_pf(p) for p in twice.paragraphs] == [_pf(p) for p in once.paragraphs]


# ───────────── синтетические случаи защиты ─────────────

def test_page_break_paragraph_not_deleted(tmp_path):
    from docx.enum.text import WD_BREAK
    doc = Document()
    doc.add_heading("ВВЕДЕНИЕ", 1)
    doc.add_paragraph("Текст первый.")
    doc.add_paragraph()
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.add_paragraph()
    doc.add_paragraph("Текст второй.")
    path = tmp_path / "pb.docx"; doc.save(path)
    fixed = Document(io.BytesIO(autofix_docx(path)))
    assert any("<w:br" in p._element.xml for p in fixed.paragraphs), "разрыв страницы удалён"


def test_list_hanging_indent_kept(tmp_path):
    doc = Document()
    doc.add_heading("ВВЕДЕНИЕ", 1)
    p = doc.add_paragraph("пункт списка", style="List Bullet")
    p.paragraph_format.first_line_indent = Cm(-0.63)
    p.paragraph_format.left_indent = Cm(1.27)
    path = tmp_path / "list.docx"; doc.save(path)
    fixed = Document(io.BytesIO(autofix_docx(path)))
    lp = [q for q in fixed.paragraphs if q.text == "пункт списка"][0]
    assert abs(lp.paragraph_format.first_line_indent.cm + 0.63) < 0.05
    assert lp.alignment is None


def test_file_extension_dot_not_glued(tmp_path):
    doc = Document()
    doc.add_heading("ВВЕДЕНИЕ", 1)
    doc.add_paragraph("Файл формата .txt хранит данные , затем читается .")
    path = tmp_path / "dot.docx"; doc.save(path)
    fixed = Document(io.BytesIO(autofix_docx(path)))
    assert fixed.paragraphs[1].text == "Файл формата .txt хранит данные, затем читается."
